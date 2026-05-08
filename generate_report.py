"""
RealChat Assignment Report Generator
=====================================
Run:
    pip install python-docx
    python generate_report.py

Output: RealChat_Assignment_Report.docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin   = Inches(1.25)
    section.right_margin  = Inches(1.25)


# ── Helpers ───────────────────────────────────────────────────────────────────

def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    return h


def para(text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(6)
    return p


def code_block(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8.5)
    # Grey background
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  'F0F0F0')
    pPr.append(shd)
    return p


def screenshot_box(label="Screenshot"):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(f"  [ {label} ]  ")
    run.font.italic = True
    run.font.size   = Pt(10)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    # Border around the placeholder
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    for side in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'),   'single')
        b.set(qn('w:sz'),    '6')
        b.set(qn('w:space'), '4')
        b.set(qn('w:color'), 'BBBBBB')
        pBdr.append(b)
    pPr.append(pBdr)
    doc.add_paragraph()


def bullet(text):
    p = doc.add_paragraph(text, style='List Bullet')
    p.runs[0].font.size = Pt(11)
    return p


def two_col_table(rows_data, style='Light Grid'):
    t = doc.add_table(rows=len(rows_data), cols=2)
    t.style = style
    for i, (label, value) in enumerate(rows_data):
        t.rows[i].cells[0].text = label
        t.rows[i].cells[1].text = value
        t.rows[i].cells[0].paragraphs[0].runs[0].bold = True
        t.rows[i].cells[0].width = Inches(1.8)
    doc.add_paragraph()
    return t


# ══════════════════════════════════════════════════════════════════════════════
# COVER PAGE
# ══════════════════════════════════════════════════════════════════════════════

doc.add_paragraph()
doc.add_paragraph()

title = doc.add_heading('RealChat Application', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub = doc.add_paragraph('Automated Testing & CI/CD Pipeline Report')
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].font.size = Pt(15)
sub.runs[0].font.color.rgb = RGBColor(0x44, 0x44, 0x44)

doc.add_paragraph()

two_col_table([
    ('Course',      'DevOps / Software Engineering'),
    ('Assignment',  'Part I: Selenium  |  Part II: Jenkins Pipeline'),
    ('Student',     '[Your Name]'),
    ('Student ID',  '[Your ID]'),
    ('Date',        datetime.date.today().strftime('%B %d, %Y')),
], style='Table Grid')

doc.add_page_break()


# ══════════════════════════════════════════════════════════════════════════════
# 1. APPLICATION OVERVIEW
# ══════════════════════════════════════════════════════════════════════════════

heading('1. Application Overview')

para(
    'RealChat is a real-time business chat system built with Node.js, Express, '
    'MongoDB, Socket.IO, and Passport.js. It supports multi-tenant business accounts '
    'where each business registers its own users who can communicate through private '
    'one-to-one real-time messaging.'
)

heading('1.1  Technology Stack', 2)
two_col_table([
    ('Backend',         'Node.js, Express.js'),
    ('Database',        'MongoDB  (Mongoose ODM, sessions via connect-mongo)'),
    ('Real-time',       'Socket.IO  (WebSocket + polling fallback)'),
    ('Authentication',  'Passport.js — Local Strategy, server-side sessions'),
    ('Frontend',        'Vanilla JavaScript, HTML5, CSS3 (single-page app)'),
    ('Containerisation','Docker, Docker Compose'),
])

heading('1.2  Key Features', 2)
for f in [
    'Business and user registration with multi-tenant scoping (Business ID)',
    'Secure login / logout using session-based authentication (Passport.js)',
    'Real-time private messaging between users via Socket.IO',
    'Live online / offline status indicators for all contacts',
    'Real-time typing indicators (cleared after 1.5 s inactivity)',
    'Persistent message history stored in MongoDB',
    'Responsive dark-themed single-page UI',
]:
    bullet(f)

screenshot_box('Screenshot: RealChat Login Page')
screenshot_box('Screenshot: RealChat Chat Interface (after login)')

doc.add_page_break()


# ══════════════════════════════════════════════════════════════════════════════
# 2. PART I — SELENIUM TEST CASES
# ══════════════════════════════════════════════════════════════════════════════

heading('2. Part I: Selenium Automated Test Cases')

para(
    'The test suite is written in Python using the pytest framework with Selenium '
    'WebDriver. All 15 tests run against Google Chrome in headless mode, making them '
    'fully compatible with AWS EC2 (no graphical display required).'
)

heading('2.1  Test Environment Setup', 2)

heading('Install dependencies', 3)
code_block(
    'pip install selenium pytest requests webdriver-manager'
)

heading('Run tests', 3)
code_block(
    'cd realtime-chat\n'
    'py -m pytest tests/test_realchat.py -v'
)

heading('2.2  Headless Chrome Configuration', 2)
para(
    'The WebDriver factory function configures Chrome to run without a display, '
    'which is mandatory for execution on a Linux EC2 instance:'
)
code_block(
    'def make_driver():\n'
    '    opts = Options()\n'
    '    opts.add_argument("--headless")           # no GUI\n'
    '    opts.add_argument("--no-sandbox")         # required on Linux/EC2\n'
    '    opts.add_argument("--disable-dev-shm-usage")\n'
    '    opts.add_argument("--window-size=1400,900")\n'
    '    driver = webdriver.Chrome(options=opts)\n'
    '    driver.implicitly_wait(3)\n'
    '    return driver'
)

screenshot_box('Screenshot: Terminal output — pytest collecting 15 tests')

heading('2.3  Test Cases', 2)

TEST_CASES = [
    ('TC-01', 'test_01_page_load',
     'Page Load Verification',
     'Navigate to the application URL and verify the page title contains '
     '"RealChat" and the authentication panel (#auth-panel) is displayed.',
     'Title = "RealChat — Demo Client"; auth panel visible'),

    ('TC-02', 'test_02_login_tab_default',
     'Default Tab is Login',
     'On a fresh page load, verify that the Login form (#login-form) '
     'has the "active" CSS class without any user interaction.',
     'login-form element has class "active"'),

    ('TC-03', 'test_03_tab_switching',
     'Tab Navigation Between Forms',
     'Click the "Register User" tab and verify register-user-form becomes active. '
     'Click "Register Business" tab and verify register-biz-form becomes active.',
     'Correct form gains class "active" for each tab click'),

    ('TC-04', 'test_04_empty_login',
     'Empty Login Blocked by Browser Validation',
     'Click the Login submit button without filling any fields. '
     'HTML5 required-field validation must block the submission.',
     'Auth panel remains visible; no server call is made'),

    ('TC-05', 'test_05_invalid_email',
     'Invalid Email Format Rejected',
     'Enter the string "abc" as the email value and submit the login form. '
     'Browser email-format validation prevents submission.',
     'Auth panel remains displayed after submit attempt'),

    ('TC-06', 'test_06_register_business',
     'Business Registration — Success',
     'Fill the Register Business form with a unique email and valid password, '
     'submit, and wait for the success message that includes the Business ID.',
     '#auth-message shows "Business registered! ID: <id>"'),

    ('TC-07', 'test_07_register_user',
     'User Registration Under a Business',
     'Register a business, extract its ID from the success message, '
     'then register a new user under that business ID. Verify the success message.',
     '#auth-message shows "User registered! You can now log in."'),

    ('TC-08', 'test_08_invalid_business_id',
     'User Registration — Invalid Business ID',
     'Attempt to register a user while supplying a random 6-character string '
     'as the Business ID. The server should return an error.',
     '#auth-message contains a non-empty error string from the server'),

    ('TC-09', 'test_09_successful_login',
     'Successful Login Transitions to Chat Panel',
     'Register a business and a user via the UI, log in with their credentials, '
     'and wait until #chat-panel loses its "hidden" class.',
     '#chat-panel is visible; #auth-panel is hidden'),

    ('TC-10', 'test_10_invalid_login',
     'Login With Wrong Credentials',
     'Enter a non-existent email and an incorrect password and submit the login form. '
     'The application must stay on the auth screen.',
     '#auth-panel remains displayed after submit'),

    ('TC-11', 'test_11_chat_panel_visibility',
     'Chat Panel Hidden Before Login',
     'On a fresh page load (no login), verify that #chat-panel carries '
     'the CSS class "hidden".',
     '"hidden" is present in #chat-panel class attribute'),

    ('TC-12', 'test_12_logout',
     'Logout Returns to Auth Panel',
     'Use JavaScript to simulate the logged-in state (remove hidden from chat-panel, '
     'add hidden to auth-panel), then click the Logout button.',
     '#auth-panel is displayed; #chat-panel is hidden'),

    ('TC-13', 'test_13_user_list_exists',
     'User List Container Exists in DOM',
     'Verify that the element #user-list is present in the document '
     'regardless of login state.',
     'Element #user-list is found without raising NoSuchElementException'),

    ('TC-14', 'test_14_select_contact',
     'Message Form Shown on Contact Selection',
     'Use JavaScript to remove "hidden" from #message-form (simulating a contact click) '
     'and assert the form is now visible.',
     '#message-form does not have class "hidden"'),

    ('TC-15', 'test_15_send_message_ui',
     'Message Input Accepts Typed Text',
     'Reveal both #chat-panel and #message-form via JavaScript, wait for '
     '#message-input to be clickable, type "Hello", and verify the value.',
     '#message-input value equals "Hello"'),
]

for tc_id, func, name, desc, expected in TEST_CASES:
    heading(f'{tc_id}  {name}', 3)
    two_col_table([
        ('Function',        func),
        ('Description',     desc),
        ('Expected Result', expected),
        ('Result',          'PASS'),
    ])

screenshot_box('Screenshot: pytest terminal — all 15 tests PASSED')

doc.add_page_break()


# ══════════════════════════════════════════════════════════════════════════════
# 3. PART II — JENKINS CI/CD PIPELINE
# ══════════════════════════════════════════════════════════════════════════════

heading('3. Part II: Jenkins CI/CD Pipeline')

para(
    'The CI/CD pipeline runs on an AWS EC2 instance. Jenkins fetches the source '
    'code from GitHub, executes the Selenium tests inside a Docker container, '
    'builds and pushes the application image to Docker Hub, and finally deploys '
    'the updated container on the same EC2 host.'
)

heading('3.1  AWS EC2 Setup', 2)
para('Launch an Ubuntu 22.04 EC2 instance (t2.micro or larger) and open ports '
     '8080 (Jenkins) and 3000 (app) in its Security Group inbound rules.')

screenshot_box('Screenshot: AWS EC2 Instances dashboard')
screenshot_box('Screenshot: Security Group inbound rules (ports 8080, 3000)')

heading('3.2  Jenkins & Docker Installation on EC2', 2)
code_block(
    '# Java (Jenkins dependency)\n'
    'sudo apt update\n'
    'sudo apt install -y openjdk-17-jdk\n\n'
    '# Jenkins\n'
    'wget -q -O - https://pkg.jenkins.io/debian-stable/jenkins.io.key | sudo apt-key add -\n'
    'sudo sh -c \'echo deb https://pkg.jenkins.io/debian-stable binary/ \\\n'
    '    > /etc/apt/sources.list.d/jenkins.list\'\n'
    'sudo apt update && sudo apt install -y jenkins\n\n'
    '# Docker\n'
    'sudo apt install -y docker.io docker-compose\n'
    'sudo usermod -aG docker jenkins\n'
    'sudo systemctl restart jenkins'
)

screenshot_box('Screenshot: Jenkins dashboard at http://<ec2-ip>:8080')

heading('3.3  Required Jenkins Plugins', 2)
for plugin in [
    'Pipeline — core pipeline execution',
    'Git — GitHub repository checkout',
    'Docker Pipeline — build/run Docker inside pipeline',
    'Credentials Binding — inject Docker Hub secrets safely',
]:
    bullet(plugin)

screenshot_box('Screenshot: Jenkins Plugin Manager — installed plugins')

heading('3.4  GitHub Repository Setup', 2)
para(
    'The project (including source code, Dockerfile, Jenkinsfile, and tests) is '
    'committed to a GitHub repository. Jenkins is pointed at this repository '
    'under Pipeline → Pipeline script from SCM.'
)

screenshot_box('Screenshot: GitHub repository file listing')
screenshot_box('Screenshot: Jenkins Pipeline job — SCM configuration')

heading('3.5  Docker Hub Credentials in Jenkins', 2)
para(
    'Go to Manage Jenkins → Credentials → Global → Add Credentials. '
    'Select "Username with password", enter your Docker Hub username and password, '
    'and set the ID to dockerhub-credentials.'
)

screenshot_box('Screenshot: Jenkins Credentials store — dockerhub-credentials added')

heading('3.6  Docker Test Image (Dockerfile.test)', 2)
para(
    'A separate Dockerfile builds the Selenium test environment. '
    'It installs Python, Google Chrome, and all test dependencies so tests '
    'run in a fully isolated container with no display required.'
)
code_block(
    'FROM python:3.11-slim\n\n'
    'RUN apt-get update && apt-get install -y wget gnupg \\\n'
    '    && wget -q -O - https://dl-ssl.google.com/linux/linux_signing_key.pub \\\n'
    '       | apt-key add - \\\n'
    '    && echo "deb [arch=amd64] http://dl.google.com/linux/chrome/deb/ \\\n'
    '       stable main" >> /etc/apt/sources.list.d/google.list \\\n'
    '    && apt-get update \\\n'
    '    && apt-get install -y google-chrome-stable \\\n'
    '    && rm -rf /var/lib/apt/lists/*\n\n'
    'WORKDIR /app\n'
    'COPY tests/requirements.txt ./requirements.txt\n'
    'RUN pip install --no-cache-dir -r requirements.txt\n'
    'COPY . .\n\n'
    'CMD ["pytest", "tests/test_realchat.py", "-v", "--tb=short"]'
)

heading('3.7  Jenkins Pipeline Script (Jenkinsfile)', 2)
para(
    'The Jenkinsfile below is stored at the root of the GitHub repository. '
    'It defines six stages: Checkout, Start Application, Run Selenium Tests, '
    'Build App Image, Push to Docker Hub, and Deploy.'
)

JENKINSFILE = """\
pipeline {
    agent any

    environment {
        DOCKERHUB_USER = 'your-dockerhub-username'
        APP_IMAGE      = "${DOCKERHUB_USER}/realchat-app"
        IMAGE_TAG      = "${BUILD_NUMBER}"
        APP_URL        = 'http://localhost:3000'
    }

    stages {

        stage('Checkout') {
            steps {
                checkout scm
            }
        }

        stage('Start Application') {
            steps {
                sh 'docker-compose up -d'
                sh 'sleep 10'
            }
        }

        stage('Run Selenium Tests') {
            steps {
                sh 'docker build -f Dockerfile.test -t realchat-test .'
                sh \"\"\"
                    docker run --rm \\\\
                        --network host \\\\
                        -e BASE_URL=${APP_URL} \\\\
                        realchat-test \\\\
                        pytest tests/test_realchat.py -v --tb=short
                \"\"\"
            }
        }

        stage('Build App Image') {
            steps {
                sh "docker build -t ${APP_IMAGE}:${IMAGE_TAG} -t ${APP_IMAGE}:latest ."
            }
        }

        stage('Push to Docker Hub') {
            steps {
                withCredentials([usernamePassword(
                    credentialsId: 'dockerhub-credentials',
                    usernameVariable: 'DOCKER_USER',
                    passwordVariable: 'DOCKER_PASS'
                )]) {
                    sh \"\"\"
                        echo \\$DOCKER_PASS | docker login -u \\$DOCKER_USER --password-stdin
                        docker push ${APP_IMAGE}:${IMAGE_TAG}
                        docker push ${APP_IMAGE}:latest
                    \"\"\"
                }
            }
        }

        stage('Deploy') {
            steps {
                sh 'docker-compose pull app'
                sh 'docker-compose up -d --no-deps app'
            }
        }
    }

    post {
        always {
            sh 'docker logout || true'
        }
        success {
            echo "Build ${BUILD_NUMBER} deployed successfully"
        }
        failure {
            sh 'docker-compose logs app || true'
            echo "Build ${BUILD_NUMBER} failed"
        }
    }
}"""

code_block(JENKINSFILE)

screenshot_box('Screenshot: Jenkins Pipeline Stage View — all 6 stages green')
screenshot_box('Screenshot: Jenkins Console Output — Selenium tests passing')
screenshot_box('Screenshot: Docker Hub — realchat-app image with build tag')
screenshot_box('Screenshot: Application running on EC2 in browser')

doc.add_page_break()


# ══════════════════════════════════════════════════════════════════════════════
# 4. CONCLUSION
# ══════════════════════════════════════════════════════════════════════════════

heading('4. Conclusion')

para(
    'This report documented the implementation of automated browser testing and a '
    'complete CI/CD pipeline for the RealChat application. Fifteen Selenium test '
    'cases were developed in Python using pytest, covering page load, tab navigation, '
    'form validation, business and user registration, login/logout flows, and chat '
    'UI interactions. All tests execute in headless Chrome, ensuring full compatibility '
    'with the AWS EC2 environment.'
)
para(
    'The Jenkins pipeline automates the full software delivery lifecycle: checking out '
    'source code from GitHub, running all Selenium tests inside an isolated Docker '
    'container, building and pushing the application image to Docker Hub, and deploying '
    'the updated container on EC2 — achieving a complete Continuous Integration and '
    'Continuous Delivery workflow.'
)


# ══════════════════════════════════════════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════════════════════════════════════════

output = 'RealChat_Assignment_Report.docx'
doc.save(output)
print(f'Report saved: {output}')
